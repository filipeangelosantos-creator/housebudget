"""Excel (.xlsx) statement reading: extract cell grid, then reuse the CSV
column-detection pipeline. Dates arrive as datetime objects, numbers as floats,
which makes detection more reliable than strings."""
import io

from openpyxl import load_workbook


def read_xlsx_rows(data: bytes) -> list[list]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    best_rows: list[list] = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else c for c in row]
            if any(str(c).strip() for c in cells):
                rows.append([c.strip() if isinstance(c, str) else c for c in cells])
        if len(rows) > len(best_rows):
            best_rows = rows
    wb.close()
    return best_rows
